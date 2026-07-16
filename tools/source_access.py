#!/usr/bin/env python3
"""D-04 deterministic source location and exact locator reopening."""

from __future__ import annotations

import hashlib
import io
import json
import logging
import math
import re
import threading
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path, PurePosixPath
from typing import Any

if __package__:
    from .evidence_registry import (
        Evidence,
        EvidenceError,
        excerpt_sha256,
        load_evidence_registry,
    )
    from .extraction_schema import (
        ExtractionSchemaError,
        LineRangeLocator,
        Locator,
        NotebookCellLocator,
        PdfPageLocator,
        SectionLocator,
        SymbolLocator,
        TableRangeLocator,
        locator_from_dict,
    )
    from .file_classification import CLASSIFICATION_SAMPLE_BYTES, classify_file
    from .notebook_extractor import (
        NotebookExtractionLimits,
        extract_notebook_bytes,
    )
    from .pdf_extractor import PdfExtractionLimits, extract_pdf_bytes
    from .project_layout import CURRENT_SCHEMA_VERSION, LayoutError
    from .project_registry import load_registered_project
    from .source_recovery import SourceRecoveryError, recover_source
    from .source_registry import SourceRecord, load_source_registry
    from .text_extractor import (
        SUPPORTED_TEXT_FORMATS,
        TextExtractionLimits,
        extract_text_bytes,
    )
else:
    from evidence_registry import (  # type: ignore[no-redef]
        Evidence,
        EvidenceError,
        excerpt_sha256,
        load_evidence_registry,
    )
    from extraction_schema import (  # type: ignore[no-redef]
        ExtractionSchemaError,
        LineRangeLocator,
        Locator,
        NotebookCellLocator,
        PdfPageLocator,
        SectionLocator,
        SymbolLocator,
        TableRangeLocator,
        locator_from_dict,
    )
    from file_classification import (  # type: ignore[no-redef]
        CLASSIFICATION_SAMPLE_BYTES,
        classify_file,
    )
    from notebook_extractor import (  # type: ignore[no-redef]
        NotebookExtractionLimits,
        extract_notebook_bytes,
    )
    from pdf_extractor import (  # type: ignore[no-redef]
        PdfExtractionLimits,
        extract_pdf_bytes,
    )
    from project_layout import (  # type: ignore[no-redef]
        CURRENT_SCHEMA_VERSION,
        LayoutError,
    )
    from project_registry import load_registered_project  # type: ignore[no-redef]
    from source_recovery import (  # type: ignore[no-redef]
        SourceRecoveryError,
        recover_source,
    )
    from source_registry import (  # type: ignore[no-redef]
        SourceRecord,
        load_source_registry,
    )
    from text_extractor import (  # type: ignore[no-redef]
        SUPPORTED_TEXT_FORMATS,
        TextExtractionLimits,
        extract_text_bytes,
    )


SOURCE_ACCESS_SCHEMA_VERSION = CURRENT_SCHEMA_VERSION
SOURCE_ACCESS_VERSION = "source-access-v1"
SOURCE_LOCATION_KIND = "llmwiki-source-location"
SOURCE_OPEN_KIND = "llmwiki-source-open-result"
TABLE_EXCERPT_FORMAT = "table-json-matrix-v1"
MAX_TABLE_RANGE_CELLS = 100_000
_MAX_SPREADSHEET_ROW = 1_048_576
_MAX_SPREADSHEET_COLUMN = 16_384
_PYPDF_LOG_LOCK = threading.RLock()

_SOURCE_ID_PATTERN = re.compile(r"src-[0-9a-f]{32}")
_EVIDENCE_ID_PATTERN = re.compile(r"evd-[0-9a-f]{64}")
_SHA256_PATTERN = re.compile(r"[0-9a-f]{64}")


class SourceAccessError(LayoutError):
    """Base error for deterministic D-04 source access."""

    reason_code = "source-access-failed"


class SourceNotFoundError(SourceAccessError):
    """Raised when a source or Evidence ID is unknown."""

    reason_code = "source-not-registered"


class SourceMissingError(SourceAccessError):
    """Raised when the current recorded path no longer resolves to a file."""

    reason_code = "source-current-path-missing"


class SourceRelocationAmbiguousError(SourceAccessError):
    """Raised when relocation has multiple equal-priority hash matches."""

    reason_code = "source-relocation-ambiguous"


class SourceBoundaryError(SourceAccessError):
    """Raised when a current path resolves outside the registered project."""

    reason_code = "source-path-outside-project"


class SourceContentMismatchError(SourceAccessError):
    """Raised when current bytes do not match the recorded source version."""

    reason_code = "source-content-hash-mismatch"


class SourceVersionMismatchError(SourceAccessError):
    """Raised when Evidence belongs to a non-current source version."""

    reason_code = "current-source-version-mismatch"


class SourceLocatorError(SourceAccessError):
    """Raised when a locator is invalid or cannot identify current content."""

    reason_code = "source-locator-invalid"


class SourceFormatError(SourceAccessError):
    """Raised when a locator/source format pair cannot be reopened."""

    reason_code = "source-locator-format-unsupported"


class SourceReadError(SourceAccessError):
    """Raised when current source bytes cannot be read exactly."""

    reason_code = "source-read-failed"


class SourceExcerptMismatchError(SourceAccessError):
    """Raised when reopened text differs from persisted Evidence."""

    reason_code = "source-excerpt-hash-mismatch"


def _is_integer(value: object) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def _source_id(value: object) -> str:
    if not isinstance(value, str) or _SOURCE_ID_PATTERN.fullmatch(value) is None:
        raise SourceNotFoundError(
            "source_id must use Core-generated form 'src-' plus 32 lowercase hex digits"
        )
    return value


def _evidence_id(value: object) -> str:
    if not isinstance(value, str) or _EVIDENCE_ID_PATTERN.fullmatch(value) is None:
        raise SourceNotFoundError(
            "evidence_id must use form 'evd-' plus 64 lowercase hex digits"
        )
    return value


def _content_hash(value: object, field_name: str) -> str:
    if not isinstance(value, str) or _SHA256_PATTERN.fullmatch(value) is None:
        raise SourceContentMismatchError(
            f"{field_name} must be a lowercase SHA-256 digest"
        )
    return value


def _excerpt_hash(value: object, field_name: str) -> str:
    if not isinstance(value, str) or _SHA256_PATTERN.fullmatch(value) is None:
        raise SourceExcerptMismatchError(
            f"{field_name} must be a lowercase SHA-256 digest"
        )
    return value


def _normalize_locator(locator: Locator | dict[str, Any]) -> Locator:
    raw: object = locator.as_dict() if isinstance(locator, Locator) else locator
    if not isinstance(raw, dict):
        raise SourceLocatorError("locator must be an object")
    version = raw.get("schema_version")
    if not _is_integer(version):
        raise SourceLocatorError("locator schema_version must be an integer")
    try:
        return locator_from_dict(raw)
    except ExtractionSchemaError as exc:
        raise SourceLocatorError(f"invalid locator: {exc}") from exc


def _reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise SourceLocatorError(f"duplicate locator JSON key {key!r}")
        result[key] = value
    return result


def _reject_json_constant(value: str) -> Any:
    raise SourceLocatorError(
        f"non-finite locator JSON constant {value!r} is not supported"
    )


def deserialize_locator(payload: str | bytes) -> Locator:
    """Strictly deserialize one C-01 locator for CLI and Core callers."""

    if isinstance(payload, bytes):
        try:
            text = payload.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise SourceLocatorError("locator JSON must be UTF-8") from exc
    elif isinstance(payload, str):
        text = payload
    else:
        raise SourceLocatorError("locator JSON must be str or bytes")
    try:
        raw = json.loads(
            text,
            object_pairs_hook=_reject_duplicate_keys,
            parse_constant=_reject_json_constant,
        )
    except json.JSONDecodeError as exc:
        raise SourceLocatorError(f"invalid locator JSON: {exc.msg}") from exc
    return _normalize_locator(raw)


@dataclass(frozen=True)
class SourceLocation:
    """Current registered path and source version, without reading content."""

    project_id: str
    source_id: str
    project_root: Path
    current_path: str
    absolute_path: Path
    current_version: int
    content_hash: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema_version": SOURCE_ACCESS_SCHEMA_VERSION,
            "kind": SOURCE_LOCATION_KIND,
            "access_version": SOURCE_ACCESS_VERSION,
            "project_id": self.project_id,
            "source_id": self.source_id,
            "project_root": str(self.project_root),
            "current_path": self.current_path,
            "absolute_path": str(self.absolute_path),
            "current_version": self.current_version,
            "content_hash": self.content_hash,
        }


@dataclass(frozen=True)
class SourceOpenResult:
    """Exact reopened excerpt after current content-hash verification."""

    source: SourceLocation
    locator: Locator
    excerpt: str
    excerpt_hash: str
    excerpt_format: str
    evidence_id: str | None = None
    evidence_source_version: int | None = None
    excerpt_hash_verified: bool | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "schema_version": SOURCE_ACCESS_SCHEMA_VERSION,
            "kind": SOURCE_OPEN_KIND,
            "access_version": SOURCE_ACCESS_VERSION,
            "source": self.source.as_dict(),
            "locator": self.locator.as_dict(),
            "excerpt": self.excerpt,
            "excerpt_hash": self.excerpt_hash,
            "excerpt_encoding": "utf-8",
            "excerpt_format": self.excerpt_format,
            "content_hash_verified": True,
            "excerpt_hash_verified": self.excerpt_hash_verified,
            "evidence_id": self.evidence_id,
            "evidence_source_version": self.evidence_source_version,
        }


def _source_record(project_id: str, source_id: str, workspace_root: str | Path) -> SourceRecord:
    normalized = _source_id(source_id)
    try:
        registry = load_source_registry(workspace_root, project_id)
    except (LayoutError, OSError) as exc:
        raise SourceAccessError(f"could not load source registry: {exc}") from exc
    try:
        return registry.by_source_id[normalized]
    except KeyError as exc:
        raise SourceNotFoundError(
            f"source_id is not registered for project {project_id}: {normalized}"
        ) from exc


def _resolve_current_path(project_root: Path, current_path: str) -> Path:
    relative = PurePosixPath(current_path)
    candidate = project_root.joinpath(*relative.parts)
    try:
        resolved = candidate.resolve(strict=True)
    except FileNotFoundError as exc:
        raise SourceMissingError(
            f"current source path does not exist: {current_path}"
        ) from exc
    except OSError as exc:
        raise SourceReadError(
            f"could not resolve current source path {current_path}: {exc}"
        ) from exc
    root = project_root.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise SourceBoundaryError(
            f"current source path resolves outside the registered project: {current_path}"
        ) from exc
    if not resolved.is_file():
        raise SourceMissingError(
            f"current source path is not a regular file: {current_path}"
        )
    return resolved


def _recover_failed_access(
    workspace_root: str | Path,
    project_id: str,
    source_id: str,
    original_error: SourceAccessError,
) -> None:
    try:
        recovery = recover_source(workspace_root, project_id, source_id)
    except SourceRecoveryError as exc:
        raise SourceAccessError(
            f"could not evaluate deterministic source relocation: {exc}"
        ) from exc
    if recovery.ambiguous:
        candidates = ", ".join(recovery.candidate_paths)
        raise SourceRelocationAmbiguousError(
            "source relocation has multiple equal-priority exact-hash candidates: "
            f"{candidates}"
        )
    if recovery.status in {"recovered", "not-needed"}:
        return
    raise original_error


def locate_source(
    workspace_root: str | Path,
    project_id: str,
    source_id: str,
) -> SourceLocation:
    """Resolve the current path, recovering identity only after access failure."""

    try:
        registration = load_registered_project(workspace_root, project_id)
    except (LayoutError, OSError) as exc:
        raise SourceAccessError(f"could not load registered project: {exc}") from exc
    source = _source_record(
        registration.project_id,
        source_id,
        workspace_root,
    )
    if source.current_version is None or source.current_content_hash is None:
        raise SourceContentMismatchError(
            f"source {source.source_id} has no recorded content version"
        )
    try:
        absolute_path = _resolve_current_path(
            registration.project_root,
            source.current_path,
        )
    except (SourceMissingError, SourceBoundaryError, SourceReadError) as exc:
        _recover_failed_access(
            workspace_root,
            registration.project_id,
            source.source_id,
            exc,
        )
        source = _source_record(
            registration.project_id,
            source.source_id,
            workspace_root,
        )
        if source.current_version is None or source.current_content_hash is None:
            raise SourceContentMismatchError(
                f"source {source.source_id} has no recorded content version"
            )
        absolute_path = _resolve_current_path(
            registration.project_root,
            source.current_path,
        )
    return SourceLocation(
        project_id=registration.project_id,
        source_id=source.source_id,
        project_root=registration.project_root,
        current_path=source.current_path,
        absolute_path=absolute_path,
        current_version=source.current_version,
        content_hash=source.current_content_hash,
    )


def _read_verified_source(location: SourceLocation) -> bytes:
    try:
        data = location.absolute_path.read_bytes()
    except OSError as exc:
        raise SourceReadError(
            f"could not read current source {location.current_path}: {exc}"
        ) from exc
    observed = hashlib.sha256(data).hexdigest()
    if observed != location.content_hash:
        raise SourceContentMismatchError(
            "current source bytes do not match the recorded source version: "
            f"expected {location.content_hash}; observed {observed}"
        )
    return data


def _line_excerpt(
    data: bytes,
    *,
    relative_path: str,
    start_line: int,
    end_line: int,
) -> str:
    try:
        classification = classify_file(
            relative_path,
            data[:CLASSIFICATION_SAMPLE_BYTES],
        )
    except Exception as exc:
        raise SourceFormatError(
            f"text source format could not be classified safely: {exc}"
        ) from exc
    format_value = classification.format
    if format_value not in SUPPORTED_TEXT_FORMATS:
        raise SourceFormatError(
            f"line locators are unsupported for source format {format_value!r}"
        )
    try:
        result = extract_text_bytes(
            data,
            relative_path=relative_path,
            format_value=format_value,
            limits=TextExtractionLimits(
                max_source_bytes=max(4, len(data)),
                max_line_characters=max(1, len(data) + 1),
                max_block_lines=max(1, len(data) + 1),
            ),
        )
    except Exception as exc:
        raise SourceFormatError(
            f"text source could not be reopened safely: {exc}"
        ) from exc
    if result.document is None or result.status not in {"processed", "partial"}:
        raise SourceFormatError(
            f"text source could not be reopened: {result.reason_code}: {result.reason}"
        )
    if any(block.truncated for block in result.document.blocks):
        raise SourceLocatorError("text reopening would return a truncated excerpt")
    text = "".join(block.text for block in result.document.blocks)
    lines = text.splitlines(keepends=True)
    if end_line > len(lines):
        raise SourceLocatorError(
            f"line range {start_line}-{end_line} exceeds {len(lines)} source lines"
        )
    return "".join(lines[start_line - 1 : end_line])


def _pdf_page_excerpt(data: bytes, *, relative_path: str, page_number: int) -> str:
    pypdf_logger = logging.getLogger("pypdf")
    try:
        with _PYPDF_LOG_LOCK:
            previous_level = pypdf_logger.level
            pypdf_logger.setLevel(max(previous_level, logging.ERROR))
            try:
                result = extract_pdf_bytes(
                    data,
                    relative_path=relative_path,
                    limits=PdfExtractionLimits(max_source_bytes=max(8, len(data))),
                )
            finally:
                pypdf_logger.setLevel(previous_level)
    except Exception as exc:
        raise SourceFormatError(
            f"PDF source could not be reopened safely: {exc}"
        ) from exc
    if result.document is None or result.status not in {"processed", "partial"}:
        raise SourceFormatError(
            f"PDF source could not be reopened: {result.reason_code}: {result.reason}"
        )
    for block in result.document.blocks:
        if isinstance(block.locator, PdfPageLocator) and (
            block.locator.page_number == page_number
        ):
            if block.truncated:
                raise SourceLocatorError("PDF page reopening would be truncated")
            return block.text
    raise SourceLocatorError(f"PDF page {page_number} does not exist")


def _notebook_cell_excerpt(
    data: bytes,
    *,
    relative_path: str,
    locator: NotebookCellLocator,
) -> str:
    try:
        result = extract_notebook_bytes(
            data,
            relative_path=relative_path,
            limits=NotebookExtractionLimits(
                max_source_bytes=max(4, len(data)),
                max_cell_characters=max(1, len(data) + 1),
            ),
        )
    except Exception as exc:
        raise SourceFormatError(
            f"Notebook source could not be reopened safely: {exc}"
        ) from exc
    if result.document is None or result.status not in {"processed", "partial"}:
        raise SourceFormatError(
            "Notebook source could not be reopened: "
            f"{result.reason_code}: {result.reason}"
        )
    source_block_id = f"cell-{locator.cell_index}-source"
    for block in result.document.blocks:
        if block.block_id != source_block_id:
            continue
        if not isinstance(block.locator, NotebookCellLocator):
            continue
        if locator.cell_id is not None and block.locator.cell_id != locator.cell_id:
            raise SourceLocatorError(
                f"Notebook cell {locator.cell_index} no longer has cell_id "
                f"{locator.cell_id!r}"
            )
        if block.truncated:
            raise SourceLocatorError("Notebook cell reopening would be truncated")
        return block.text
    raise SourceLocatorError(f"Notebook cell {locator.cell_index} does not exist")


def _table_json_value(value: object) -> object:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise SourceLocatorError("table range contains a non-finite number")
        return value
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {
            "type": "timedelta",
            "microseconds": (
                value.days * 86_400_000_000
                + value.seconds * 1_000_000
                + value.microseconds
            ),
        }
    if isinstance(value, Decimal):
        return {"type": "decimal", "value": str(value)}
    raise SourceLocatorError(
        f"table range contains unsupported cell value type {type(value).__name__}"
    )


def _table_range_excerpt(data: bytes, *, locator: TableRangeLocator) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:
        raise SourceFormatError(
            "openpyxl is required to reopen spreadsheet table ranges"
        ) from exc
    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise SourceFormatError(f"spreadsheet could not be opened safely: {exc}") from exc
    try:
        if locator.sheet not in workbook.sheetnames:
            raise SourceLocatorError(
                f"spreadsheet sheet does not exist: {locator.sheet!r}"
            )
        worksheet = workbook[locator.sheet]
        min_column, min_row, max_column, max_row = range_boundaries(
            f"{locator.start_cell}:{locator.end_cell}"
        )
        if max_row > _MAX_SPREADSHEET_ROW or max_column > _MAX_SPREADSHEET_COLUMN:
            raise SourceLocatorError(
                "table range exceeds OOXML worksheet row or column bounds"
            )
        cell_count = (max_row - min_row + 1) * (max_column - min_column + 1)
        if cell_count > MAX_TABLE_RANGE_CELLS:
            raise SourceLocatorError(
                "table range exceeds deterministic cell limit: "
                f"{cell_count} cells requested; limit is {MAX_TABLE_RANGE_CELLS}"
            )
        try:
            matrix = [
                [_table_json_value(cell.value) for cell in row]
                for row in worksheet.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_column,
                    max_col=max_column,
                )
            ]
        except SourceAccessError:
            raise
        except Exception as exc:
            raise SourceFormatError(
                f"spreadsheet cells could not be reopened safely: {exc}"
            ) from exc
        try:
            return json.dumps(
                matrix,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
                allow_nan=False,
            )
        except (TypeError, ValueError) as exc:
            raise SourceLocatorError(
                "table range cannot be serialized canonically"
            ) from exc
    finally:
        workbook.close()


def _reopen_locator(
    data: bytes,
    *,
    relative_path: str,
    locator: Locator,
) -> tuple[str, str]:
    if isinstance(locator, LineRangeLocator):
        return (
            _line_excerpt(
                data,
                relative_path=relative_path,
                start_line=locator.start_line,
                end_line=locator.end_line,
            ),
            "text-line-range",
        )
    if isinstance(locator, (SectionLocator, SymbolLocator)):
        return (
            _line_excerpt(
                data,
                relative_path=relative_path,
                start_line=locator.start_line,
                end_line=locator.end_line,
            ),
            "text-line-range",
        )
    if isinstance(locator, PdfPageLocator):
        return (
            _pdf_page_excerpt(
                data,
                relative_path=relative_path,
                page_number=locator.page_number,
            ),
            "pdf-extracted-page-text",
        )
    if isinstance(locator, NotebookCellLocator):
        return (
            _notebook_cell_excerpt(
                data,
                relative_path=relative_path,
                locator=locator,
            ),
            "notebook-cell-source",
        )
    if isinstance(locator, TableRangeLocator):
        return _table_range_excerpt(data, locator=locator), TABLE_EXCERPT_FORMAT
    raise SourceFormatError(
        f"unsupported locator class {type(locator).__name__}"
    )


def open_source(
    workspace_root: str | Path,
    project_id: str,
    *,
    source_id: str,
    locator: Locator | dict[str, Any],
    expected_content_hash: str | None = None,
    expected_excerpt_hash: str | None = None,
    evidence_id: str | None = None,
    evidence_source_version: int | None = None,
) -> SourceOpenResult:
    """Open one locator at the current path and verify exact current bytes."""

    normalized_locator = _normalize_locator(locator)
    location = locate_source(workspace_root, project_id, source_id)
    def verify_requested_identity(candidate: SourceLocation) -> None:
        if evidence_source_version is not None:
            if not _is_integer(evidence_source_version) or evidence_source_version < 1:
                raise SourceVersionMismatchError(
                    "evidence_source_version must be a positive integer"
                )
            if evidence_source_version != candidate.current_version:
                raise SourceVersionMismatchError(
                    "Evidence source version is not the current registered version: "
                    f"evidence {evidence_source_version}; "
                    f"current {candidate.current_version}"
                )
        if expected_content_hash is not None:
            expected_content = _content_hash(
                expected_content_hash,
                "expected_content_hash",
            )
            if expected_content != candidate.content_hash:
                raise SourceContentMismatchError(
                    "requested content hash is not the current recorded source version: "
                    f"expected {expected_content}; current {candidate.content_hash}"
                )

    verify_requested_identity(location)
    try:
        data = _read_verified_source(location)
    except (SourceContentMismatchError, SourceReadError) as exc:
        _recover_failed_access(
            workspace_root,
            location.project_id,
            location.source_id,
            exc,
        )
        location = locate_source(
            workspace_root,
            location.project_id,
            location.source_id,
        )
        verify_requested_identity(location)
        data = _read_verified_source(location)
    excerpt, excerpt_format = _reopen_locator(
        data,
        relative_path=location.current_path,
        locator=normalized_locator,
    )
    try:
        digest = excerpt_sha256(excerpt)
    except EvidenceError as exc:
        raise SourceFormatError(
            f"reopened excerpt cannot use the strict UTF-8 Evidence contract: {exc}"
        ) from exc
    excerpt_verified: bool | None = None
    if expected_excerpt_hash is not None:
        expected_excerpt = _excerpt_hash(
            expected_excerpt_hash,
            "expected_excerpt_hash",
        )
        if digest != expected_excerpt:
            raise SourceExcerptMismatchError(
                "reopened excerpt does not match expected excerpt hash: "
                f"expected {expected_excerpt}; observed {digest}"
            )
        excerpt_verified = True
    return SourceOpenResult(
        source=location,
        locator=normalized_locator,
        excerpt=excerpt,
        excerpt_hash=digest,
        excerpt_format=excerpt_format,
        evidence_id=evidence_id,
        evidence_source_version=evidence_source_version,
        excerpt_hash_verified=excerpt_verified,
    )


def open_evidence(
    workspace_root: str | Path,
    project_id: str,
    evidence_id: str,
) -> SourceOpenResult:
    """Reopen and verify one persisted Evidence against the current source."""

    normalized_id = _evidence_id(evidence_id)
    try:
        registry = load_evidence_registry(workspace_root, project_id)
    except (LayoutError, OSError) as exc:
        raise SourceAccessError(f"could not load Evidence registry: {exc}") from exc
    try:
        evidence: Evidence = registry.by_evidence_id[normalized_id]
    except KeyError as exc:
        raise SourceNotFoundError(
            f"evidence_id is not registered for project {project_id}: {normalized_id}"
        ) from exc
    return open_source(
        workspace_root,
        project_id,
        source_id=evidence.source_id,
        locator=evidence.locator,
        expected_content_hash=evidence.content_hash,
        expected_excerpt_hash=evidence.excerpt_hash,
        evidence_id=evidence.evidence_id,
        evidence_source_version=evidence.source_version,
    )
